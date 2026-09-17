#!/usr/bin/env python3
"""
Known-answer tests for scripts/migrate_comment_extracts.py (the writer).

Self-contained (no pytest needed):  python3 scripts/test_migrate_comment_extracts.py

The writer makes no judgement about coding: the plan it is handed comes from the R
reader. What it MUST get right is never touching the workbook it reads, never
overwriting anything, and refusing rather than clobbering pruning work that has
already been done by hand.
"""

import importlib.util
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

_spec = importlib.util.spec_from_file_location(
    "mce", str(Path(__file__).with_name("migrate_comment_extracts.py")))
mce = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mce)

_passed = _failed = 0


def check(cond, msg):
    global _passed, _failed
    if cond:
        _passed += 1
        print("  ok   " + msg)
    else:
        _failed += 1
        print("  FAIL " + msg)


def fixture(tmp):
    """A one-sheet appendix and a plan for it."""
    wb_path = Path(tmp) / "appendix.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Engagement"
    ws.append(["ID", "Noteworthy", "Comment", "Overall Sentiment", "Pay"])
    ws.append([1, None, "a comment", "3", "3"])
    wb.save(wb_path)
    plan_path = Path(tmp) / "plan.json"
    plan = [{"sheet": "Engagement Extracts", "base": "Engagement",
             "rows": [{"id": "1", "themes": "Pay; Workload", "text": "a fragment"}]}]
    plan_path.write_text(json.dumps(plan), encoding="utf-8")
    return wb_path, plan_path


def main():
    with tempfile.TemporaryDirectory() as tmp:
        wb_path, plan_path = fixture(tmp)
        before = wb_path.read_bytes()
        out = Path(tmp) / "migrated.xlsx"

        added = mce.write_extracts(str(wb_path), str(plan_path), str(out))
        check(added == [("Engagement Extracts", 1)], "reports what it wrote")
        check(wb_path.read_bytes() == before,
              "the workbook it read is byte-for-byte untouched")

        wb = openpyxl.load_workbook(out)
        check(wb.sheetnames == ["Engagement", "Engagement Extracts"],
              "the extracts sheet is added beside the question sheet")
        ws = wb["Engagement Extracts"]
        rows = list(ws.iter_rows(values_only=True))
        check(rows[0] == ("ID", "Theme", "Extract", "Lead"), "headers are the four columns")
        check(rows[1] == ("1", "Pay; Workload", "a fragment", None),
              "the row carries id, every coded theme, the text, and a blank Lead")
        check(ws.freeze_panes == "A2", "the header row is frozen for the analyst")
        # The original question sheet is carried through unchanged.
        check([c for c in wb["Engagement"].iter_rows(values_only=True)][1][2] == "a comment",
              "the question sheet's own rows survive the copy")

        # Refuses to overwrite, and leaves the existing file alone.
        stamp = out.read_bytes()
        code = subprocess.run(
            [sys.executable, str(Path(__file__).with_name("migrate_comment_extracts.py")),
             "--workbook", str(wb_path), "--plan", str(plan_path), "--out", str(out)],
            capture_output=True, text=True)
        check(code.returncode != 0 and "never overwrites" in (code.stdout + code.stderr),
              "refuses to overwrite an existing output")
        check(out.read_bytes() == stamp, "and the existing output is untouched")

        # Refuses when the workbook already carries that extracts sheet, rather than
        # clobbering pruning that has been done by hand.
        wb2 = openpyxl.load_workbook(wb_path)
        wb2.create_sheet("Engagement Extracts")
        wb2.save(wb_path)
        out2 = Path(tmp) / "again.xlsx"
        code2 = subprocess.run(
            [sys.executable, str(Path(__file__).with_name("migrate_comment_extracts.py")),
             "--workbook", str(wb_path), "--plan", str(plan_path), "--out", str(out2)],
            capture_output=True, text=True)
        check(code2.returncode != 0, "refuses a workbook that already has the sheet")
        check(not out2.exists(), "and leaves no half-written file behind")

    print(f"\n{_passed} passed, {_failed} failed")
    return 1 if _failed else 0


if __name__ == "__main__":
    sys.exit(main())
