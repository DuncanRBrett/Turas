#!/usr/bin/env python3
"""
Write the extracts sheets proposed by migrate_comment_extracts.R.

The classification (which column is the verbatim, which are themes) is the R
reader's, passed in as a plan, so this script makes no judgement about the coding.
It copies the workbook, appends one "<Sheet> Extracts" worksheet per plan entry,
and saves under a NEW name. It never opens the original for writing.

openpyxl rather than openxlsx: an openxlsx load-and-save round trip collapses each
sheet's declared dimension, which openpyxl tools then choke on
(docs/HANDOVER_openxlsx_broken_workbooks.md).

Usage (called by the R script, or directly):
  python3 scripts/migrate_comment_extracts.py --workbook IN.xlsx --plan PLAN.json --out OUT.xlsx
"""
import argparse
import json
import os
import shutil
import sys

import openpyxl

HEADERS = ["ID", "Theme", "Extract", "Lead"]


def write_extracts(workbook, plan_path, out_path):
    """Copy `workbook` to `out_path` and append the planned extracts sheets."""
    if os.path.exists(out_path):
        sys.exit(f"Refusing: {out_path} already exists. This script never overwrites.")
    with open(plan_path, encoding="utf-8") as fh:
        plan = json.load(fh)

    shutil.copyfile(workbook, out_path)
    wb = openpyxl.load_workbook(out_path)
    added = []
    for entry in plan:
        name = entry["sheet"]
        if name in wb.sheetnames:
            os.remove(out_path)
            sys.exit(f"Refusing: {workbook} already has a sheet called '{name}'.")
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        for row in entry.get("rows", []):
            # Lead is left blank: which fragment leads a slide is the analyst's call.
            ws.append([row["id"], row["themes"], row["text"], None])
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 44
        ws.column_dimensions["C"].width = 90
        ws.column_dimensions["D"].width = 7
        ws.freeze_panes = "A2"
        added.append((name, len(entry.get("rows", []))))
    wb.save(out_path)
    return added


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--plan", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    added = write_extracts(args.workbook, args.plan, args.out)
    for name, n in added:
        print(f"  wrote {name}: {n} row(s)")


if __name__ == "__main__":
    main()
